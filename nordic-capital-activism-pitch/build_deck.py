#!/usr/bin/env python3
"""Builds the PowerPoint deck: Nordic Capital Engaged Equities strategy pitch."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x0B, 0x25, 0x45)
NAVY2 = RGBColor(0x13, 0x3A, 0x63)
GOLD = RGBColor(0xC4, 0x9A, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x55, 0x5F, 0x6E)
LIGHT = RGBColor(0xF2, 0xF4, 0xF8)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_rect(slide, x, y, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, runs, size=14, color=GREY, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.0, font="Georgia"):
    """runs: str, or list of (text, dict-of-overrides) paragraphs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, (text, ov) in enumerate(runs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = ov.get("align", align)
        para.line_spacing = ov.get("line_spacing", line_spacing)
        para.space_after = Pt(ov.get("space_after", 4))
        r = para.add_run()
        r.text = text
        f = r.font
        f.name = ov.get("font", font)
        f.size = Pt(ov.get("size", size))
        f.bold = ov.get("bold", bold)
        f.italic = ov.get("italic", False)
        f.color.rgb = ov.get("color", color)
    return tb


def bullet_box(slide, x, y, w, h, items, size=13, gap=8, color=RGBColor(0x2B, 0x33, 0x40)):
    """items: list of str or (bold_lead, rest) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(gap)
        para.line_spacing = 1.05
        b = para.add_run()
        b.text = "▪  "
        b.font.size = Pt(size)
        b.font.color.rgb = GOLD
        b.font.name = "Georgia"
        if isinstance(it, tuple):
            lead, rest = it
            r1 = para.add_run()
            r1.text = lead
            r1.font.bold = True
            r1.font.size = Pt(size)
            r1.font.color.rgb = NAVY
            r1.font.name = "Georgia"
            r2 = para.add_run()
            r2.text = rest
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
            r2.font.name = "Georgia"
        else:
            r = para.add_run()
            r.text = it
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.name = "Georgia"
    return tb


def new_slide(title, kicker=None):
    slide = prs.slides.add_slide(BLANK)
    add_rect(slide, 0, 0, SW, Inches(0.95), NAVY)
    add_rect(slide, 0, Inches(0.95), SW, Pt(3), GOLD)
    if kicker:
        add_text(slide, Inches(0.55), Inches(0.08), Inches(11), Inches(0.3),
                 kicker.upper(), size=10, color=GOLD, bold=True)
        ty = Inches(0.34)
    else:
        ty = Inches(0.18)
    add_text(slide, Inches(0.55), ty, Inches(12.2), Inches(0.6), title,
             size=24, color=WHITE, bold=True)
    add_text(slide, Inches(0.55), Inches(7.12), Inches(12.2), Inches(0.3),
             "Strictly private & confidential — illustrative discussion material",
             size=8, color=RGBColor(0x9A, 0xA4, 0xB2))
    return slide


def card(slide, x, y, w, h, title, body, title_size=13, body_size=11.5):
    add_rect(slide, x, y, w, h, LIGHT)
    add_rect(slide, x, y, Pt(3.5), h, GOLD)
    add_text(slide, x + Inches(0.18), y + Inches(0.10), w - Inches(0.3), Inches(0.45),
             title, size=title_size, color=NAVY, bold=True)
    add_text(slide, x + Inches(0.18), y + Inches(0.52), w - Inches(0.3), h - Inches(0.6),
             body, size=body_size, color=GREY, line_spacing=1.05)


def stat(slide, x, y, w, value, label, value_size=30):
    add_text(slide, x, y, w, Inches(0.55), value, size=value_size, color=GOLD, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, x, y + Inches(0.58), w, Inches(0.65), label, size=11, color=WHITE,
             align=PP_ALIGN.CENTER, line_spacing=1.0)


def table_slide(slide, x, y, w, headers, rows, col_widths, row_h=0.5, size=11):
    from pptx.util import Inches as In
    t = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w,
                               In(row_h * (len(rows) + 1))).table
    for j, cw in enumerate(col_widths):
        t.columns[j].width = In(cw)
    for j, htext in enumerate(headers):
        cell = t.cell(0, j)
        cell.text = htext
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size = Pt(size)
        para.runs[0].font.bold = True
        para.runs[0].font.color.rgb = WHITE
        para.runs[0].font.name = "Georgia"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
            for para in cell.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(size - 0.5)
                    r.font.color.rgb = RGBColor(0x2B, 0x33, 0x40)
                    r.font.name = "Georgia"
    return t


# ============ 1. Title ============
slide = prs.slides.add_slide(BLANK)
add_rect(slide, 0, 0, SW, SH, NAVY)
add_rect(slide, 0, Inches(3.62), SW, Pt(3), GOLD)
add_text(slide, Inches(1), Inches(1.45), Inches(11.3), Inches(0.4),
         "STRICTLY PRIVATE & CONFIDENTIAL", size=12, color=GOLD, bold=True,
         align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(2.0), Inches(11.3), Inches(1.0),
         "Nordic Capital Engaged Equities", size=44, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.0), Inches(11.3), Inches(0.5),
         "A constructive activism strategy for US and European public markets",
         size=19, color=RGBColor(0xC9, 0xD3, 0xE0), align=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(4.05), Inches(11.3), Inches(0.4),
         "Proposal for a USD 1.5bn engaged public-equities vehicle",
         size=14, color=GOLD, align=PP_ALIGN.CENTER, bold=True)
add_text(slide, Inches(1), Inches(5.9), Inches(11.3), Inches(0.8),
         [("Prepared for the Managing Partner & Executive Committee", {"align": PP_ALIGN.CENTER, "size": 12, "color": RGBColor(0xC9, 0xD3, 0xE0)}),
          ("June 2026  |  Company references are based on public information only", {"align": PP_ALIGN.CENTER, "size": 10, "color": RGBColor(0x8A, 0x96, 0xA6)})])

# ============ 2. Executive summary ============
slide = new_slide("Executive summary", "The proposal")
bullet_box(slide, Inches(0.55), Inches(1.25), Inches(7.4), Inches(5.6), [
    ("Launch “Nordic Capital Engaged Equities”: ", "a USD 1.5bn (cap USD 2.0bn) "
     "constructive-activism fund taking 2–10% stakes in 6–8 listed US and European "
     "companies (USD 2–15bn market cap) in the firm's core sectors."),
    ("Apply the firm's PE playbook in public markets: ", "operational improvement, "
     "portfolio separation, capital-allocation reform — via nomination committees and "
     "board seats in Europe, universal-proxy-backed engagement in the US."),
    ("The window is open: ", "record global activism with an M&A-heavy mix that favours "
     "a PE sponsor, 30–50% value gaps in de-rated quality mid-caps on both continents — "
     "and no practitioner combining operating depth with a transatlantic mandate."),
    ("Nordic Capital's edge is structural: ", "operating credibility, the best director "
     "bench in Europe, a New York platform, and sector depth in healthcare, tech & "
     "payments, financial services and industrials. US campaign strategy led by Knut "
     "Nyman, joining from Elliott."),
    ("Platform synergies compound: ", "take-private and PIPE pipeline for the buyout "
     "funds, carve-out deal flow, and a new fee stream — accretive from year two."),
], size=13.5, gap=12)
add_rect(slide, Inches(8.35), Inches(1.25), Inches(4.4), Inches(5.6), NAVY)
stat(slide, Inches(8.5), Inches(1.55), Inches(4.1), "USD 1.5bn", "Target size (hard cap USD 2.0bn)")
stat(slide, Inches(8.5), Inches(2.85), Inches(4.1), "6–8", "Core positions · ~50/50 US & Europe · 2–10% stakes")
stat(slide, Inches(8.5), Inches(4.15), Inches(4.1), "15%+", "Target net IRR over a cycle")
stat(slide, Inches(8.5), Inches(5.45), Inches(4.1), "H2 2027", "First close & initial deployment")

# ============ 3. Why now ============
slide = new_slide("Why now: the opening for a transatlantic engagement franchise", "Opportunity")
card(slide, Inches(0.55), Inches(1.3), Inches(6.1), Inches(1.7), "Record activism, M&A-centric",
     "2025 set a record for global activist campaigns; ~6 in 10 late-2025 campaigns carried "
     "M&A demands (break-ups, divestitures, sale processes) — terrain where a private equity "
     "sponsor has advantages no hedge-fund activist can match.")
card(slide, Inches(6.85), Inches(1.3), Inches(6.0), Inches(1.7), "Value gaps on both continents",
     "Europe: structural discount to the US, compounded in under-covered mid-caps. US: the "
     "post-2021 de-rating left fallen-angel quality in medtech, payments and software at "
     "30–50% discounts — below the multiples our buyout funds pay in auctions.")
card(slide, Inches(0.55), Inches(3.2), Inches(6.1), Inches(1.7), "Mis-shaped competition",
     "US activism is crowded at the mega-cap end but thin in operationally complex mid-caps; "
     "in Europe, Cevian has drifted large-cap and has no US book. Nobody combines PE operating "
     "depth with a transatlantic mandate.")
card(slide, Inches(6.85), Inches(3.2), Inches(6.0), Inches(1.7), "Supportive shareholder base",
     "Index funds and pension institutions on both continents increasingly back well-argued "
     "engagement agendas — lowering the cost of building winning coalitions.")
add_rect(slide, Inches(0.55), Inches(5.15), Inches(12.3), Inches(1.55), NAVY)
add_text(slide, Inches(0.85), Inches(5.38), Inches(11.8), Inches(1.2),
         [("The strategic point", {"size": 12, "color": GOLD, "bold": True}),
          ("The category-defining sponsor-backed engagement franchise spanning the US and "
           "Europe has not yet been built. Nordic Capital can build it — with advantages "
           "neither Cevian nor the US funds have: a buyout platform behind the engagement and "
           "an operating bench measured in hundreds of professionals.",
           {"size": 14, "color": WHITE, "line_spacing": 1.1})])

# ============ 4. Nordic governance ============
slide = new_slide("Two engagement regimes, one playbook", "Opportunity")
table_slide(slide, Inches(0.55), Inches(1.35), Inches(12.3),
            ["Dimension", "Europe / Nordics (home field)", "United States"],
            [
                ["Path to the board",
                 "Shareholder-elected nomination committees (Nordics: top 3–4 holders as of "
                 "Aug/Sep records); sphere engagement; 4–8% wins influence without a proxy "
                 "fight",
                 "Universal proxy card (2022) slashed the cost of board challenges; most "
                 "campaigns settle for seats; 2–6% stakes suffice with a credible slate"],
                ["Escalation tools",
                 "Low EGM thresholds (10% in Sweden); public shareholder registers enable "
                 "precise coalition mapping",
                 "13D platform; precision proxy contests; settlement negotiation; deep "
                 "advisor ecosystem"],
                ["Culture & cost",
                 "Consensus boards; private argumentation usually sufficient; low campaign "
                 "cost",
                 "Faster, transactional; settlements within 6–12 months; higher cost, "
                 "higher liquidity"],
                ["Our edge",
                 "Domestic legitimacy of a 30-year Stockholm-rooted owner; Europe's best "
                 "director bench; Cevian-validated model, no second practitioner",
                 "Sector operating credibility US activists lack; PE-grade separation "
                 "expertise; campaign leadership hired from the US activist front line"],
            ],
            [1.8, 5.3, 5.2], row_h=0.95, size=11)
add_text(slide, Inches(0.55), Inches(5.85), Inches(12.3), Inches(1.0),
         [("Same underwriting, same value plans, different influence mechanics. Timing "
           "discipline in both: Q3 ownership snapshots seat Nordic nomination committees; "
           "US nomination windows gate the proxy season.",
           {"size": 12.5, "color": GREY, "italic": True, "line_spacing": 1.15})])

# ============ 5. Why Nordic Capital ============
slide = new_slide("Why Nordic Capital wins", "Right to win")
cards = [
    ("Operating credibility", "30 years and >EUR 31bn AUM built on operational "
     "transformation. When we tell a board the margin gap is 400bps, we speak as an owner "
     "that has closed it repeatedly in private portfolios."),
    ("The director bench", "An unmatched network of Nordic chairs, former CEOs and "
     "operating advisors to propose as board candidates — the resource that decides "
     "outcomes in a nomination-committee system."),
    ("Legitimacy + US credibility", "European spheres and institutions trust a "
     "Stockholm-rooted owner; in the US, our New York platform and a campaign team hired "
     "from the activist front line give standing European funds never built."),
    ("Sector underwriting edge", "Healthcare, tech & payments, financial services, "
     "industrials — the same lens in Boston, Amsterdam and Stockholm; a universe of "
     "businesses we have studied, bid on or competed against."),
    ("Two-way PE pipeline", "Engagement surfaces take-privates and PIPEs for the flagship "
     "funds; activist-catalysed carve-outs create proprietary deal flow (strict info "
     "barriers; see governance slide)."),
    ("LP demand", "Existing LPs want Nordic exposure with liquidity; Evolution II's "
     "four-month hard-cap close shows the franchise's fundraising velocity."),
]
for i, (t, b) in enumerate(cards):
    x = Inches(0.55 + (i % 3) * 4.18)
    y = Inches(1.3 + (i // 3) * 2.75)
    card(slide, x, y, Inches(3.95), Inches(2.55), t, b, body_size=11.5)

# ============ 6. Strategy ============
slide = new_slide("Investment strategy: constructive activism, board-led", "Strategy")
add_text(slide, Inches(0.55), Inches(1.2), Inches(12.3), Inches(0.85),
         [("Take 2–10% positions in 6–8 listed US and European companies (USD 2–15bn market "
           "cap) with ≥30% upside to intrinsic value, and close the gap in 2–4 years through "
           "actions the company itself can take — driven from the boardroom under each "
           "market's governance regime.",
           {"size": 14.5, "color": NAVY, "bold": True, "line_spacing": 1.15})])
steps = [
    ("1. Underwrite", "M 0–4", "PE-grade diligence; operational benchmark; SOTP; red-team; "
     "≥30% upside with limited structural downside"),
    ("2. Accumulate", "M 3–9", "2–10% via market + blocks, timed to each market's board "
     "calendar; chair meeting before crossing 5%"),
    ("3. Engage", "M 6–18", "Private value plan to board & anchors; EU: nomination-committee "
     "seat; US: negotiated refreshment under universal-proxy optionality"),
    ("4. Execute", "M 12–36", "Margin program / separation / balance-sheet reset; quarterly "
     "milestones vs the underwrite"),
    ("5. Realise", "M 24–48", "Re-rating, strategic sale — or conflict-cleared take-private "
     "by our buyout funds"),
]
for i, (t, m, b) in enumerate(steps):
    x = Inches(0.55 + i * 2.52)
    add_rect(slide, x, Inches(2.35), Inches(2.32), Inches(2.9), NAVY if i % 2 == 0 else NAVY2)
    add_text(slide, x + Inches(0.12), Inches(2.5), Inches(2.1), Inches(0.4), t,
             size=13.5, color=GOLD, bold=True)
    add_text(slide, x + Inches(0.12), Inches(2.92), Inches(2.1), Inches(0.3), m,
             size=10, color=RGBColor(0x9A, 0xA4, 0xB2), bold=True)
    add_text(slide, x + Inches(0.12), Inches(3.25), Inches(2.1), Inches(1.9), b,
             size=10.5, color=WHITE, line_spacing=1.1)
bullet_box(slide, Inches(0.55), Inches(5.55), Inches(12.3), Inches(1.4), [
    ("Doctrine: ", "constructive-first — public escalation needs IC + reputation-committee "
     "sign-off; never a hostile opening move; no shorting of engagement targets; max 2 "
     "concurrent public-phase campaigns."),
], size=12.5)

# ============ 7. Thesis archetypes ============
slide = new_slide("Three thesis archetypes", "Strategy")
arch = [
    ("Margin-gap operators", "40–50% of book",
     "Quality franchises running 300–600bps below achievable margins (execution, supply "
     "chain, quality systems — not market position). Value plan: board-sponsored "
     "operational program with hard milestones.", "Examples: Elekta, Smith+Nephew"),
    ("Conglomerate / SOTP discounts", "30–40% of book",
     "Multi-division companies where separation, divestiture or monetisation of hidden "
     "assets unlocks 30–60% upside. Plays to the record M&A-activism mix and our carve-out "
     "expertise.", "Examples: Philips, Baxter"),
    ("Capital allocation & strategic catalysts", "15–25% of book",
     "Lazy balance sheets, value-destructive M&A patterns, or companies that should start "
     "or accept strategic processes; often overlaps the other archetypes.",
     "Examples: Global Payments, Zimmer Biomet"),
]
for i, (t, w, b, ex) in enumerate(arch):
    x = Inches(0.55 + i * 4.18)
    add_rect(slide, x, Inches(1.35), Inches(3.95), Inches(4.7), LIGHT)
    add_rect(slide, x, Inches(1.35), Inches(3.95), Pt(4), GOLD)
    add_text(slide, x + Inches(0.2), Inches(1.55), Inches(3.55), Inches(0.75), t,
             size=15.5, color=NAVY, bold=True)
    add_text(slide, x + Inches(0.2), Inches(2.3), Inches(3.55), Inches(0.35), w,
             size=11.5, color=GOLD, bold=True)
    add_text(slide, x + Inches(0.2), Inches(2.7), Inches(3.55), Inches(2.4), b,
             size=12, color=GREY, line_spacing=1.15)
    add_text(slide, x + Inches(0.2), Inches(5.35), Inches(3.55), Inches(0.6), ex,
             size=11, color=NAVY, bold=True)
add_text(slide, Inches(0.55), Inches(6.35), Inches(12.3), Inches(0.5),
         [("Sector focus mirrors the PE franchise: Healthcare & MedTech 30–40% · Industrials & "
           "Business Services 25–35% · Tech, Software & Payments 15–25% · Financial Services "
           "10–20%. Geography: US 40–50% · Europe 50–60% (of which Nordics 20–30%).",
           {"size": 11.5, "color": GREY, "italic": True})])

# ============ 8. Funnel ============
slide = new_slide("Target universe: a disciplined funnel", "Universe")
funnel = [
    ("~1,300", "US & European listed names in our four core sectors, USD 2–15bn, "
     "post liquidity screens (~55% US / 45% EU)", 12.3),
    ("~300", "Quality screen: defensible position, sound gross margins, clear peer set", 10.2),
    ("~80–100", "Value gap ≥30% on normalised margins or SOTP", 8.1),
    ("~40", "Path to influence: open register, no hostile blocker, no defeating "
     "defence structures", 6.0),
    ("12–16", "Live underwritten pipeline (after MNPI / conflicts clearance)", 4.2),
    ("6–8", "Funded core positions", 2.6),
]
y = 1.3
for i, (n, label, w) in enumerate(funnel):
    x = Inches(0.55 + (12.3 - w) / 2)
    shade = [NAVY, NAVY2, RGBColor(0x1E, 0x4D, 0x7B), RGBColor(0x2E, 0x60, 0x91),
             RGBColor(0x44, 0x74, 0xA4), GOLD][i]
    add_rect(slide, x, Inches(y), Inches(w), Inches(0.72), shade)
    add_text(slide, x + Inches(0.25), Inches(y + 0.08), Inches(1.5), Inches(0.5), n,
             size=18, color=WHITE if i < 5 else NAVY, bold=True)
    add_text(slide, x + Inches(1.8), Inches(y + 0.14), Inches(w - 2.0), Inches(0.5), label,
             size=12.5, color=WHITE if i < 5 else NAVY)
    y += 0.88
add_text(slide, Inches(0.55), Inches(6.7), Inches(12.3), Inches(0.5),
         [("USD 1.5bn across 6–8 cores = USD 180–280m per position = 4–10% of a USD 3–6bn "
           "European mid-cap or 2–5% of a USD 5–15bn US name — influential under both "
           "regimes, in the band where competition is thinnest.",
           {"size": 12.5, "color": NAVY, "bold": True})])

# ============ 9–11. Pipeline ============
slide = new_slide("Illustrative pipeline (1/2) — public information only", "Pipeline")
pipe1 = [
    ("Elekta  ·  Sweden  ·  MedTech", "Margin-gap operator",
     "Global #2 in radiation oncology in a growing oligopoly; heavily de-rated, with "
     "third-party fair-value estimates implying ~50%+ discounts during 2026. FY25/26 adj. "
     "EBIT margin ~12% vs mid-to-high teens at Siemens Healthineers/Varian.",
     "Board-sponsored margin program; order-to-revenue execution reset; service "
     "acceleration; strategic optionality. Home-field nomination-committee engagement."),
    ("Smith+Nephew  ·  UK  ·  MedTech", "Margin gap & structural options",
     "Quality ortho/sports-med/wound franchise lagging peers on margin and execution for "
     "years; engagement validated (Cevian on the register since 2024) with the turnaround "
     "only partly delivered.",
     "Hold the board to peer-level margins; evaluate separation of underperforming "
     "franchises and a US listing review — a second credible engaged owner can be "
     "decisive."),
    ("Philips  ·  Netherlands  ·  HealthTech", "Recovery & SOTP",
     "Emerging from the Respironics recall era (US litigation settled) with margins well "
     "below pre-crisis levels and peers; anchor shareholder (Exor) shows the register is "
     "open to engaged owners.",
     "Margin normalisation with board accountability; portfolio review of Personal Health "
     "vs the healthcare core — force the SOTP into the price."),
]
y = 1.3
for name, arch_t, sit, plan in pipe1:
    add_rect(slide, Inches(0.55), Inches(y), Inches(12.3), Inches(1.78), LIGHT)
    add_rect(slide, Inches(0.55), Inches(y), Pt(4), Inches(1.78), GOLD)
    add_text(slide, Inches(0.8), Inches(y + 0.08), Inches(5.2), Inches(0.35), name,
             size=13.5, color=NAVY, bold=True)
    add_text(slide, Inches(8.6), Inches(y + 0.08), Inches(4.1), Inches(0.35), arch_t,
             size=11, color=GOLD, bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(0.8), Inches(y + 0.45), Inches(6.6), Inches(1.3),
             [("Situation: " + sit, {"size": 10.5, "color": GREY, "line_spacing": 1.06})])
    add_text(slide, Inches(7.6), Inches(y + 0.45), Inches(5.1), Inches(1.3),
             [("Value plan: " + plan, {"size": 10.5, "color": RGBColor(0x2B, 0x33, 0x40),
                                       "line_spacing": 1.06})])
    y += 1.92

slide = new_slide("Illustrative pipeline (2/2) — public information only", "Pipeline")
pipe2 = [
    ("Baxter International  ·  US  ·  MedTech", "Completing the simplification",
     "Multi-year simplification in motion (BioPharma Solutions sold; Vantive kidney-care "
     "divested to Carlyle, early 2025), yet the equity kept de-rating amid execution "
     "stumbles and a CEO transition — a fallen angel far below peer multiples.",
     "Finish the job: complete the separation program, reset the cost base, refresh the "
     "board, disciplined capital-return framework. Classic universal-proxy-era engagement."),
    ("Global Payments  ·  US  ·  Payments", "Capital allocation & credibility reset",
     "Scaled merchant-acquiring franchise at a deeply depressed multiple after the poorly "
     "received 2025 Worldpay / Issuer Solutions restructuring; activist interest publicly "
     "reported — the value gap is recognised.",
     "Integration milestones with board accountability; divest non-core; binding "
     "capital-return commitments post-deleveraging; payments operating expertise to the "
     "board (our Nets/Nexi heritage)."),
    ("Zimmer Biomet  ·  US  ·  MedTech", "Capital-allocation catalyst",
     "Orthopaedics leader compounding modestly at a structurally depressed multiple; "
     "investor confidence eroded by guidance resets and M&A choices.",
     "Capital-allocation framework (organic + buybacks over dilutive M&A); prune subscale "
     "adjacencies; margin program toward peer levels."),
]
y = 1.3
for name, arch_t, sit, plan in pipe2:
    add_rect(slide, Inches(0.55), Inches(y), Inches(12.3), Inches(1.5), LIGHT)
    add_rect(slide, Inches(0.55), Inches(y), Pt(4), Inches(1.5), GOLD)
    add_text(slide, Inches(0.8), Inches(y + 0.06), Inches(5.6), Inches(0.32), name,
             size=13, color=NAVY, bold=True)
    add_text(slide, Inches(8.6), Inches(y + 0.06), Inches(4.1), Inches(0.32), arch_t,
             size=10.5, color=GOLD, bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(0.8), Inches(y + 0.40), Inches(6.6), Inches(1.05),
             [("Situation: " + sit, {"size": 9.8, "color": GREY, "line_spacing": 1.03})])
    add_text(slide, Inches(7.6), Inches(y + 0.40), Inches(5.1), Inches(1.05),
             [("Value plan: " + plan, {"size": 9.8, "color": RGBColor(0x2B, 0x33, 0x40),
                                       "line_spacing": 1.03})])
    y += 1.64
add_rect(slide, Inches(0.55), Inches(6.25), Inches(12.3), Inches(0.85), NAVY)
add_text(slide, Inches(0.85), Inches(6.33), Inches(11.8), Inches(0.7),
         [("Watchlist (survives screens 1–3)", {"size": 10.5, "color": GOLD, "bold": True,
                                                "space_after": 1}),
          ("EU: Getinge · GN Store Nord · Orkla · Tieto · Demant · Securitas · ISS · Valmet · "
           "Storebrand   |   US: Teleflex · Henry Schein · Fiserv · specialty insurers",
           {"size": 11, "color": WHITE, "line_spacing": 1.05})])

# ============ 12. Fund terms ============
slide = new_slide("Vehicle, size & terms", "Structure")
table_slide(slide, Inches(0.55), Inches(1.3), Inches(12.3),
            ["Parameter", "Recommendation"],
            [
                ["Target size", "USD 1.5bn target; USD 2.0bn hard cap — sized to the USD 2–15bn "
                 "transatlantic mid-cap band where influence is cheap and competition thin"],
                ["Structure", "Luxembourg RAIF (AIFMD); evergreen with 3-year soft lock, then "
                 "quarterly liquidity with 25% investor-level gates"],
                ["GP commitment", "EUR 150–200m from balance sheet + partners/employees"],
                ["Fees", "1.25–1.50% management; 17.5% performance over max(50/50 S&P 500 / "
                 "MSCI Europe + 200bps, 5% absolute); 3-year crystallisation; high-water mark"],
                ["Concentration", "Max 25% NAV per position at cost (35% at market); max 2 "
                 "concurrent public-escalation campaigns"],
                ["Co-invest", "Sleeves for oversized situations and take-private bridges"],
                ["Target returns", "15%+ net IRR over a cycle; catalyst-driven, entry "
                 "discounts ≥30% provide downside margin"],
            ],
            [2.4, 9.9], row_h=0.68, size=12)

# ============ 12b. Hedged or unhedged ============
slide = new_slide("Run it hedged or unhedged?", "Structure")
hedge_opts = [
    ("Long-only, unhedged", "The Cevian model",
     "Net ~100%; returns = beta + engagement alpha. Simple, cheap, matches LP "
     "expectations for engaged equity — but fully exposed in risk-off years, mid-campaign.",
     False),
    ("Long + tactical overlay", "RECOMMENDED",
     "Net 80–100% by default. IC may hedge up to ~30% of NAV via liquid index "
     "futures/options under pre-defined stress triggers. FX systematically hedged to "
     "USD. Keeps the risk premium and catalyst optionality; caps tail drawdowns cheaply.",
     True),
    ("Market-neutral activist", "Rejected",
     "Permanent shorts isolate campaign alpha at 150–300bps annual drag — compounding "
     "against multi-year horizons; basis risk in a 6–8 name book; a structural short "
     "book sits awkwardly with a constructive board-seat brand.",
     False),
]
for i, (t, tag, b, rec) in enumerate(hedge_opts):
    x = Inches(0.55 + i * 4.18)
    add_rect(slide, x, Inches(1.3), Inches(3.95), Inches(3.4), NAVY if rec else LIGHT)
    add_rect(slide, x, Inches(1.3), Inches(3.95), Pt(4), GOLD)
    add_text(slide, x + Inches(0.2), Inches(1.5), Inches(3.55), Inches(0.7), t,
             size=15.5, color=GOLD if rec else NAVY, bold=True)
    add_text(slide, x + Inches(0.2), Inches(2.2), Inches(3.55), Inches(0.35), tag,
             size=11, color=(RGBColor(0x9A, 0xA4, 0xB2) if not rec else WHITE), bold=True)
    add_text(slide, x + Inches(0.2), Inches(2.6), Inches(3.55), Inches(2.0), b,
             size=11.5, color=WHITE if rec else GREY, line_spacing=1.12)
bullet_box(slide, Inches(0.55), Inches(5.0), Inches(12.3), Inches(1.8), [
    ("Why: ", "the alpha is idiosyncratic and catalyst-driven; entry discounts ≥30% are "
     "the primary downside protection. A permanent hedge premium would cost ~2–3% p.a. "
     "to neutralise the beta LPs are deliberately buying — protection the evergreen "
     "structure (3-year lock, gates) already provides against forced selling."),
    ("Hard rule: ", "no single-name shorts, ever — not on targets, not on peers "
     "(conflicts, disclosure, reputation). Hedge P&L stays inside the performance-fee "
     "benchmark; the 50/50 S&P 500 / MSCI Europe hurdle is unchanged."),
], size=12, gap=6)

# ============ 13. Team ============
slide = new_slide("Team & organisation: small, senior, transatlantic, conflict-proofed", "Team")
add_rect(slide, Inches(4.3), Inches(1.2), Inches(4.7), Inches(0.72), NAVY)
add_text(slide, Inches(4.3), Inches(1.29), Inches(4.7), Inches(0.5),
         "Head of Engaged Equities (Partner)\nEx-Cevian / top activist franchise · ExCo · Stockholm",
         size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
# Knut Nyman — highlighted founding role
add_rect(slide, Inches(0.55), Inches(2.08), Inches(12.3), Inches(0.95), NAVY2)
add_rect(slide, Inches(0.55), Inches(2.08), Pt(4), Inches(0.95), GOLD)
add_text(slide, Inches(0.85), Inches(2.16), Inches(11.8), Inches(0.35),
         [("Knut Nyman — Principal, US Situations & Campaign Strategy  ·  New York  ·  "
           "joining from Elliott Investment Management",
           {"size": 12.5, "color": GOLD, "bold": True})])
add_text(slide, Inches(0.85), Inches(2.52), Inches(11.8), Inches(0.45),
         [("Owns the US book: idea generation & underwriting · 13D and universal-proxy "
           "playbook · settlement negotiation · US advisor network. Co-leads 1–2 campaigns "
           "per year; campaign lead from year two; path to Partner; carry from day one.",
           {"size": 10.5, "color": WHITE, "line_spacing": 1.05})])
boxes = [
    ("Deputy PM / Head of Research", "Underwriting standards · red-team owner"),
    ("Sector Principal — Healthcare/MedTech", "Campaign lead (one per region)"),
    ("Sector Principal — Industrials/Tech", "Campaign lead (one per region)"),
    ("4–5 Investment professionals", "Modelling · register analytics · split SHB/NY"),
    ("Head of Stewardship & Governance", "Nomination committees · director sourcing · proxy"),
    ("Dedicated Compliance Officer", "Owns the info barrier · reports to Group GC"),
]
for i, (t, b) in enumerate(boxes):
    x = Inches(0.55 + (i % 3) * 4.18)
    y = Inches(3.18 + (i // 3) * 1.12)
    add_rect(slide, x, y, Inches(3.95), Inches(0.98), LIGHT)
    add_rect(slide, x, y, Pt(3.5), Inches(0.98), GOLD)
    add_text(slide, x + Inches(0.15), y + Inches(0.06), Inches(3.7), Inches(0.4), t,
             size=11, color=NAVY, bold=True)
    add_text(slide, x + Inches(0.15), y + Inches(0.5), Inches(3.7), Inches(0.42), b,
             size=9.5, color=GREY)
bullet_box(slide, Inches(0.55), Inches(5.55), Inches(12.3), Inches(1.5), [
    ("Conflicts framework (make-or-break): ", "hard information barrier vs the PE business; "
     "restricted-list screening before any accumulation; wall-crossing only via documented "
     "GC approval."),
    ("Take-private protocol: ", "pre-agreed, LPAC-disclosed terms (independent pricing, "
     "roll-or-cash optionality, conflicted-party abstentions) if a position becomes a "
     "buyout target."),
    ("Reputation committee: ", "Managing Partner sign-off for any public escalation, media "
     "engagement or EGM/proxy action."),
], size=10.5, gap=4)

# ============ 14. Economics & risks ============
slide = new_slide("Platform economics, risks & mitigants", "Economics & risk")
add_text(slide, Inches(0.55), Inches(1.2), Inches(12.3), Inches(0.35),
         "Economics for the platform", size=14, color=NAVY, bold=True)
bullet_box(slide, Inches(0.55), Inches(1.6), Inches(12.3), Inches(1.2), [
    "~USD 20m annual management fees at target size vs ~USD 12–14m steady-state costs — "
    "accretive from year two; ~USD 30–35m performance fees per cycle at 15% gross.",
    "Strategic P&L is bigger: one incremental take-private or carve-out sourced through "
    "the strategy likely exceeds the vehicle's annual fee income in platform value.",
], size=12, gap=5)
table_slide(slide, Inches(0.55), Inches(2.85), Inches(12.3),
            ["Risk", "Mitigant"],
            [
                ["Reputational spillover to the PE brand",
                 "Constructive-first doctrine; reputation-committee veto; no hostile openings"],
                ["Conflicts / MNPI vs private strategies",
                 "Hard info barrier; dedicated compliance; LPAC-disclosed take-private protocol"],
                ["Sphere control blocks change",
                 "Screened out ex ante; engage spheres as allies where present"],
                ["European mid-cap liquidity (thinner than US lines)",
                 "ADV-tied sizing; larger tickets in US names; evergreen + gates; co-invest "
                 "sleeves; event-driven exits"],
                ["Key-person dependence",
                 "Co-PM structure; institutional IC; management-company equity"],
                ["Talent competition (Active Value Partners et al.)",
                 "Move in 2026; platform economics and brand beat start-up offers"],
            ],
            [4.6, 7.7], row_h=0.55, size=11)

# ============ 15. Roadmap ============
slide = new_slide("Implementation roadmap", "Execution")
phases = [
    ("Phase 0 — Decide & design", "Jul–Sep 2026",
     "ExCo approval · conflicts framework with GC · regulatory scoping (Lux RAIF / AIFM "
     "passport) · launch Head search"),
    ("Phase 1 — Build", "Oct 2026 – Mar 2027",
     "Head + Deputy hired · Knut Nyman & NY pod on board · core team of 7 · EUR 150–200m "
     "seed closed · 12–16 underwritten names · LP pre-marketing"),
    ("Phase 2 — First close & deploy", "Apr–Sep 2027",
     "First close ~USD 750m · 2–3 positions built against the board calendar (Q3-2027 "
     "Nordic snapshots; US nomination windows for the 2028 proxy season)"),
    ("Phase 3 — Scale", "Oct 2027 – Dec 2028",
     "Final close USD 1.5–2.0bn · 6–8 positions across both regions · first board seats "
     "and settlements in the 2028 cycles · first realisations 2029"),
]
for i, (t, when, b) in enumerate(phases):
    x = Inches(0.55 + i * 3.14)
    add_rect(slide, x, Inches(1.5), Inches(2.95), Inches(3.6), NAVY if i % 2 == 0 else NAVY2)
    add_rect(slide, x, Inches(1.5), Inches(2.95), Pt(4), GOLD)
    add_text(slide, x + Inches(0.15), Inches(1.7), Inches(2.65), Inches(0.8), t,
             size=13.5, color=GOLD, bold=True)
    add_text(slide, x + Inches(0.15), Inches(2.5), Inches(2.65), Inches(0.35), when,
             size=11, color=RGBColor(0x9A, 0xA4, 0xB2), bold=True)
    add_text(slide, x + Inches(0.15), Inches(2.9), Inches(2.65), Inches(2.1), b,
             size=11, color=WHITE, line_spacing=1.15)
add_rect(slide, Inches(0.55), Inches(5.5), Inches(12.3), Inches(1.2), LIGHT)
add_text(slide, Inches(0.85), Inches(5.68), Inches(11.8), Inches(0.9),
         [("The ask", {"size": 12, "color": GOLD, "bold": True}),
          ("Executive Committee mandate to proceed with Phase 0: conflicts framework, "
           "regulatory scoping and the Head of Engaged Equities search — decision point for "
           "full launch in September 2026.", {"size": 13.5, "color": NAVY, "bold": True})])

# ============ 16. Closing ============
slide = prs.slides.add_slide(BLANK)
add_rect(slide, 0, 0, SW, SH, NAVY)
add_rect(slide, 0, Inches(3.1), SW, Pt(3), GOLD)
add_text(slide, Inches(1.2), Inches(2.0), Inches(11), Inches(1.0),
         "The transatlantic engagement franchise has not been built yet.",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1.2), Inches(3.35), Inches(11), Inches(0.6),
         "Nordic Capital is the firm to build it.", size=22, color=GOLD, bold=True,
         align=PP_ALIGN.CENTER)
add_text(slide, Inches(1.2), Inches(6.5), Inches(11), Inches(0.6),
         "Illustrative discussion material · Company references based on public information "
         "only · Not investment advice", size=9, color=RGBColor(0x8A, 0x96, 0xA6),
         align=PP_ALIGN.CENTER)

# ============ A1. Appendix: Cevian record ============
slide = new_slide("Appendix — Precedents (1/3): Cevian, the constructive model at scale",
                  "Appendix")
table_slide(slide, Inches(0.55), Inches(1.2), Inches(12.3),
            ["Campaign", "Period", "What happened", "Result"],
            [
                ["Volvo Group", "2006–17", "Portfolio focus, cost discipline; stake sold to "
                 "Geely for ~USD 3.9bn — reported among its most profitable", "Win"],
                ["ThyssenKrupp", "2013–c.22", "Break-up thesis; Elevator sale (EUR 17bn) "
                 "vindicated SOTP but fragile balance sheet destroyed equity en route",
                 "Failure"],
                ["RSA Insurance", "2013–21", "Focus & disposals, then sale; Intact/Tryg "
                 "takeover at ~50% premium", "Win"],
                ["ABB", "2015–25", "Power Grids exit to Hitachi, decentralisation, "
                 "buybacks; long hold sold down 2025 after a strong re-rating", "Win"],
                ["Ericsson", "2017–c.23", "Chair change, cost focus; doubled into 2021, "
                 "compliance issues gave much back", "Mixed+"],
                ["Panalpina", "c.2013–19", "Long-held ~12%; pushed sale; DSV takeover "
                 "2019 at a substantial premium", "Win"],
                ["Pearson", "2020–", "Backed digital turnaround; Apollo approaches "
                 "rejected above-market; stake since lifted to ~18% — largest disclosed",
                 "Win, ongoing"],
                ["Aviva", "2021–c.24", "Demanded ~GBP 5bn capital return; delivered, "
                 "re-rated", "Win"],
                ["UBS", "2023–25", "Post-CS re-rating bet; reported ~2x on early tranches "
                 "at 2025 sales", "Win"],
                ["Baloise", "2023–25", "Strategy reset, vote caps removed; Helvetia merger "
                 "2025; sold ~CHF 184 vs ~CHF 130 entry (~+40%)", "Win"],
                ["Smith+Nephew", "2024–", "Entered ~5%, lifted above 10% by mid-2026 "
                 "(>USD 1.3bn); margin recovery, structural options incl. US listing "
                 "review", "Ongoing"],
                ["Akzo Nobel", "2025–", "Built ~5% pressing for reset; backed the USD "
                 "25bn Axalta merger (Nov 2025), doubled to ~10.2%", "Ongoing"],
                ["SIG Group", "2025–", "New stake in the Swiss packaging maker; de-rated "
                 "quality, self-help thesis", "Ongoing"],
            ],
            [1.7, 1.0, 8.4, 1.2], row_h=0.4, size=9)
add_text(slide, Inches(0.55), Inches(6.85), Inches(12.3), Inches(0.35),
         [("Also: Skandia, Danske Bank, Tieto, Cookson/Vesuvius–Alent, Metso, "
           "Wolseley/Ferguson, Autoliv/Veoneer, Nordea (wins); Bilfinger, Vodafone "
           "(failures). ~3 clear failures in ~23 campaigns; reported low-to-mid-teens net "
           "annualised since 2002; wins cluster where a transaction crystallised value — "
           "our M&A archetype.",
           {"size": 9.5, "color": GREY, "italic": True, "line_spacing": 1.0})])

# ============ A2. Appendix: sponsor precedents ============
slide = new_slide("Appendix — Precedents (2/3): sponsors entering public-market activism",
                  "Appendix")
card(slide, Inches(0.55), Inches(1.2), Inches(6.1), Inches(2.35),
     "EQT Public Value (2018–24) — the cautionary tale",
     "Listed Nordic mid-caps: Securitas, BHG, Storebrand, BioGaia, AFRY, ~10% of Storytel "
     "(Sep 2021, became largest shareholder). Never reached scale, no board-seat "
     "playbook, growth-momentum entries; the 2022 de-rating crushed the book and the fund "
     "was liquidated — the Storytel stake sold Aug 2024.", body_size=10.5)
card(slide, Inches(6.85), Inches(1.2), Inches(6.0), Inches(2.35),
     "KKR – Henry Schein (2024–26) — the template",
     "Ananym campaigns (Nov 2024) → KKR USD 250m investment + 2 board seats (Jan 2025) → "
     "~12% completed May 2025 → CEO succession: Lowery, ex-Thermo Fisher (Mar 2026) → "
     "KKR's Dan Daniel becomes Independent Chairman (May 2026), KKR at 16.4%. Zero to "
     "largest active holder + chairmanship in ~15 months, in our core sector.",
     body_size=10.5)
add_rect(slide, Inches(0.55), Inches(3.7), Inches(12.3), Inches(1.35), LIGHT)
add_rect(slide, Inches(0.55), Inches(3.7), Pt(4), Inches(1.35), GOLD)
add_text(slide, Inches(0.85), Inches(3.78), Inches(11.8), Inches(0.3),
         [("More sponsor entries — “occasional activists” are a defining 2025–26 trend",
           {"size": 11.5, "color": NAVY, "bold": True})])
add_text(slide, Inches(0.85), Inches(4.1), Inches(11.8), Inches(0.9),
         [("KKR–US Foods (2020): USD 500m PIPE + board seat at the COVID trough, strong "
           "reported exit  ·  Silver Lake–Expedia/Twitter/Airbnb (2020): crisis PIPEs with "
           "board representation  ·  Apollo–Western Digital (2023): USD 900m convertible "
           "riding Elliott's separation push (SanDisk spin completed 2025)  ·  Warburg "
           "Pincus–ESR (2022–25): minority toehold converted into co-leading the ~USD 7bn "
           "take-private  ·  Triton–Caverion (2022–23): stake-build → tender battle → "
           "take-private at a large premium",
           {"size": 10.5, "color": GREY, "line_spacing": 1.15})])
add_rect(slide, Inches(0.55), Inches(5.25), Inches(12.3), Inches(1.55), NAVY)
add_text(slide, Inches(0.85), Inches(5.35), Inches(11.8), Inches(0.3),
         [("What the precedents are designed into", {"size": 11.5, "color": GOLD,
                                                     "bold": True})])
add_text(slide, Inches(0.85), Inches(5.68), Inches(11.8), Inches(1.05),
         [("Scale + seed from day one — not EQT PV's subscale drift  ·  Partner-level "
           "activist leadership and a US campaign Principal — not PE generalists  ·  "
           "Stakes sized for influence — not passive minorities  ·  ≥30% value-gap "
           "entries — not growth momentum  ·  Evergreen capital — never forced to sell at "
           "the bottom  ·  Screens exclude the Cevian failure fingerprint: fragile "
           "balance sheets, political stakeholders, catalysts outside shareholders' "
           "control.",
           {"size": 11, "color": WHITE, "line_spacing": 1.15})])

# ============ A3. EQT Public Value deep-dive + Irenic / Ananym ============
slide = new_slide("Appendix — Precedents (3/3): the model that fails vs the models that work",
                  "Appendix")

# Left column: EQT Public Value case study
add_rect(slide, Inches(0.55), Inches(1.2), Inches(6.05), Inches(5.5), LIGHT)
add_rect(slide, Inches(0.55), Inches(1.2), Pt(4), Inches(5.5), RGBColor(0xB0, 0x3A, 0x2E))
add_text(slide, Inches(0.75), Inches(1.32), Inches(5.7), Inches(0.4),
         "EQT Public Value (2018–24) — what NOT to do", size=14, color=NAVY, bold=True)
add_text(slide, Inches(0.75), Inches(1.78), Inches(5.7), Inches(0.95),
         [("The closest analogue to this proposal — a top-tier Nordic PE sponsor in listed "
           "mid-caps — and it failed on design, not concept. Book drifted from value (Securitas, "
           "Storebrand) to momentum (BHG, Storytel: ~10% bought Sep 2021, ~80% de-rating, "
           "force-sold Aug 2024 on the fund clock).",
           {"size": 10.5, "color": GREY, "line_spacing": 1.08})])
flaws = [
    ("Sub-scale capital", "Influence-sized stakes (2–10%) + EUR 150–200m seed"),
    ("No activist leadership", "Partner Head + US Campaign Principal (Knut Nyman)"),
    ("Quality/momentum entries", "≥30% value-gap with a named catalyst; no momentum"),
    ("Finite fund clock", "Evergreen + 3-yr lock — never a forced seller"),
]
yy = 2.85
add_text(slide, Inches(0.75), yy - 0.18, Inches(5.7), Inches(0.25),
         [("Four flaws → this proposal's answers", {"size": 10.5, "color": NAVY, "bold": True})])
for flaw, fix in flaws:
    add_text(slide, Inches(0.75), Inches(yy + 0.05), Inches(5.7), Inches(0.7),
             [("✗ " + flaw, {"size": 10, "color": RGBColor(0xB0, 0x3A, 0x2E), "bold": True,
                             "space_after": 1}),
              ("→ " + fix, {"size": 10, "color": RGBColor(0x1E, 0x6B, 0x3A)})])
    yy += 0.74
add_text(slide, Inches(0.75), Inches(yy + 0.0), Inches(5.7), Inches(0.55),
         [("Lesson: a sponsor's public vehicle only works if built as an engagement platform "
           "from day one — not a long-only fund with a famous logo.",
           {"size": 10, "color": NAVY, "italic": True, "line_spacing": 1.05})])

# Right column: Irenic & Ananym
add_text(slide, Inches(6.85), Inches(1.2), Inches(6.0), Inches(0.35),
         [("New-generation strategic activists — the model that WORKS",
           {"size": 14, "color": NAVY, "bold": True})])
add_text(slide, Inches(6.85), Inches(1.58), Inches(6.0), Inches(0.5),
         [("Concentrated separation/sale theses (our SOTP archetype), constructive-first, "
           "run by Elliott / JANA / Engine No. 1 alumni — the talent pool this team draws on.",
           {"size": 10, "color": GREY, "line_spacing": 1.05})])
add_text(slide, Inches(6.85), Inches(2.15), Inches(6.0), Inches(0.3),
         [("Irenic Capital (est. 2021, Adam Katz ex-Elliott) — ~$335m → ~$2.5bn AUM; "
           "+14/19/17% '23–'25", {"size": 10.5, "color": GOLD, "bold": True})])
irenic = [
    ["Barnes Group", "Aerospace SOTP → sale", "Board seats; sold to Apollo"],
    ["News Corp", "Split real estate from media", "Special-committee review"],
    ["Integer Holdings", "Board refresh + sale", "2 seats; strategic review '26"],
    ["Atkore / Workiva", "Sale / ops + 2 seats", "Ongoing"],
    ["Reservoir, Snap, HPE", "Strategic review / costs", "Ongoing"],
]
table_slide(slide, Inches(6.85), Inches(2.5), Inches(6.0),
            ["Target", "Thesis", "Outcome"], irenic, [1.6, 2.5, 1.9],
            row_h=0.36, size=8.5)
add_text(slide, Inches(6.85), Inches(4.75), Inches(6.0), Inches(0.3),
         [("Ananym Capital (est. 2024, Penner ex-Engine No. 1/Exxon + Silver) — ~$260m, "
           "~10 names", {"size": 10.5, "color": GOLD, "bold": True})])
ananym = [
    ["Henry Schein", "Costs, CEO, board (≤6 seats)", "Catalysed KKR deal"],
    ["Baker Hughes", "Spin off oilfield services", "Ongoing (≥60% upside)"],
    ["Siemens Energy", "Separate wind (Gamesa)", "Ongoing"],
    ["LKQ", "Divest Europe", "Ongoing"],
]
table_slide(slide, Inches(6.85), Inches(5.1), Inches(6.0),
            ["Target", "Thesis", "Outcome"], ananym, [1.6, 2.5, 1.9],
            row_h=0.36, size=8.5)
add_text(slide, Inches(6.85), Inches(6.95), Inches(6.0), Inches(0.3),
         [("Edge vs them: they have the craft but no operating platform, no EU "
           "nomination-committee home field, no buyout balance sheet for take-privates.",
           {"size": 9, "color": NAVY, "italic": True})])

prs.save("/home/user/knutnyman/nordic-capital-activism-pitch/Nordic_Capital_Engaged_Equities_Deck.pptx")
print("Deck saved with", len(prs.slides.__iter__.__self__._sldIdLst), "slides.")
